"""Build the local, indexed Census 2011 PCA SQLite database from the supplied XLSX."""

from __future__ import annotations

import argparse
import sqlite3
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.services.demographics.census_2011 import normalize


REQUIRED = {"State", "District", "Subdistt", "Town/Village", "Level", "Name", "TRU", "No_HH", "TOT_P", "TOT_M", "TOT_F", "P_06", "P_LIT", "TOT_WORK_P"}


def parse_number(value: object) -> int | None:
    if value is None or str(value).strip() in {"", "NA", "N.A."}:
        return None
    return int(value)


def build(source: Path, destination: Path) -> None:
    workbook = load_workbook(source, read_only=True, data_only=True)
    worksheet = workbook["Data"]
    headers = tuple(next(worksheet.values))
    if not REQUIRED.issubset(headers):
        missing = sorted(REQUIRED - set(headers))
        raise ValueError(f"Workbook Data sheet is missing required columns: {', '.join(missing)}")
    column = {name: headers.index(name) for name in headers}
    temporary = destination.with_suffix(destination.suffix + ".importing")
    if temporary.exists():
        temporary.unlink()
    destination.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(temporary)
    db.executescript("""
        PRAGMA journal_mode=OFF;
        PRAGMA synchronous=OFF;
        CREATE TABLE census_geo (
          state_code TEXT NOT NULL, district_code TEXT NOT NULL, subdistrict_code TEXT NOT NULL, village_code TEXT NOT NULL,
          level TEXT NOT NULL, name TEXT NOT NULL, name_key TEXT NOT NULL, state_key TEXT NOT NULL, district_key TEXT NOT NULL,
          tru TEXT NOT NULL, households INTEGER, total_population INTEGER, male_population INTEGER, female_population INTEGER,
          population_0_6 INTEGER, literate_population INTEGER, working_population INTEGER,
          PRIMARY KEY (state_code, district_code, subdistrict_code, village_code, level, tru)
        );
    """)
    insert = "INSERT INTO census_geo VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"
    batch = []
    count = 0
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        value = lambda key: row[column[key]]
        state_code, district_code, subdistrict_code, village_code = (str(value(key)).zfill(width) for key, width in (("State", 2), ("District", 3), ("Subdistt", 5), ("Town/Village", 6)))
        name = str(value("Name")).strip()
        batch.append((state_code, district_code, subdistrict_code, village_code, str(value("Level")).upper(), name, normalize(name), normalize(state_code and row[column["Name"]] if str(value("Level")).upper() == "STATE" else ""), "", str(value("TRU")), parse_number(value("No_HH")), parse_number(value("TOT_P")), parse_number(value("TOT_M")), parse_number(value("TOT_F")), parse_number(value("P_06")), parse_number(value("P_LIT")), parse_number(value("TOT_WORK_P"))))
        # Derive parent name keys from code maps below after importing, avoiding assumptions about row order.
        if len(batch) >= 5000:
            db.executemany(insert, batch); count += len(batch); batch.clear()
    if batch:
        db.executemany(insert, batch); count += len(batch)
    db.executescript("""
        CREATE TABLE state_names AS SELECT state_code, name AS state_name, name_key AS state_key FROM census_geo WHERE level='STATE' AND tru='Total';
        CREATE TABLE district_names AS SELECT state_code, district_code, name AS district_name, name_key AS district_key FROM census_geo WHERE level='DISTRICT' AND tru='Total';
        UPDATE census_geo SET state_key=COALESCE((SELECT state_key FROM state_names s WHERE s.state_code=census_geo.state_code), state_key);
        UPDATE census_geo SET district_key=COALESCE((SELECT district_key FROM district_names d WHERE d.state_code=census_geo.state_code AND d.district_code=census_geo.district_code), district_key);
        CREATE INDEX idx_census_district ON census_geo(level, tru, state_key, name);
        CREATE INDEX idx_census_subdistrict ON census_geo(level, tru, state_key, district_key, name);
        CREATE INDEX idx_census_snapshot ON census_geo(level, state_key, district_key, name_key, tru);
    """)
    db.commit(); db.close(); workbook.close()
    temporary.replace(destination)
    print(f"Imported {count:,} Census 2011 PCA records into {destination}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parents[1] / "data" / "census_2011.sqlite")
    arguments = parser.parse_args()
    build(arguments.source, arguments.output)
